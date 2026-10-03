"""Shared encoder-only NLP pipeline for the document review assistant.

The application intentionally uses three BERT model roles:
1. BERT fine-tuned for Named Entity Recognition (NER).
2. BERT fine-tuned for extractive Question Answering (QA).
3. Base BERT embeddings for ambiguity, summary, and clause review.

No generative model is used. Outputs are evidence from the source document or
scores based on BERT embeddings.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Iterable

import fitz
import torch
from docx import Document
from transformers import AutoModel, AutoModelForQuestionAnswering, AutoModelForTokenClassification, AutoTokenizer, pipeline


NER_MODEL_NAME = "dslim/bert-base-NER"
QA_MODEL_NAME = "deepset/bert-base-cased-squad2"
EMBEDDING_MODEL_NAME = "bert-base-uncased"

# NER guesses below this score are mostly noise, such as "and" tagged as an organisation
MIN_ENTITY_SCORE = 0.6

CLAUSE_DESCRIPTIONS = {
    "Confidentiality": "confidential information must be protected and not disclosed",
    "Payment": "invoice payment fee price and payment deadline obligation",
    "Termination": "ending the agreement cancellation notice and termination rights",
    "Data handling": "personal data information sharing security storage and privacy",
    "Liability": "damages indemnity liability responsibility and limitation of loss",
}


@dataclass
class Finding:
    """A detected entity or structured sensitive-data finding."""

    text: str
    category: str
    confidence: float | None = None


def extract_text(file_name: str, file_bytes: bytes) -> str:
    """Extract text from a TXT, PDF, or DOCX file uploaded by the user."""
    suffix = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if suffix == "txt":
        return file_bytes.decode("utf-8", errors="replace")
    if suffix == "pdf":
        pdf = fitz.open(stream=file_bytes, filetype="pdf")
        return "\n".join(page.get_text() for page in pdf)
    if suffix == "docx":
        document = Document(BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_rows: list[str] = []
        for table_number, table in enumerate(document.tables, start=1):
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_lines = [re.sub(r"\s+", " ", line).strip() for line in cell.text.splitlines() if line.strip()]
                    cells.append(" ; ".join(cell_lines))
                # Merged DOCX cells can appear more than once in a row.
                unique_cells = list(dict.fromkeys(cell for cell in cells if cell))
                if unique_cells:
                    table_rows.append(f"Table {table_number}: " + " | ".join(unique_cells))
        return "\n".join(paragraphs + table_rows)

    raise ValueError("Please upload a TXT, PDF, or DOCX document.")


def normalise_text(text: str) -> str:
    """Remove repeated whitespace while keeping paragraph boundaries readable."""
    paragraphs = [re.sub(r"\s+", " ", paragraph).strip() for paragraph in text.splitlines()]
    return "\n".join(paragraph for paragraph in paragraphs if paragraph)


def split_sentences(text: str) -> list[str]:
    """Use lightweight sentence splitting so no extra language model is needed."""
    cleaned = re.sub(r"\s+", " ", text).strip()
    return [sentence.strip() for sentence in re.split(r"(?<=[.!?])\s+", cleaned) if sentence.strip()]


def split_for_qa(text: str, chunk_words: int = 230, overlap_words: int = 35) -> list[str]:
    """Create overlapping chunks that fit within BERT QA's token context window."""
    words = text.split()
    if not words:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = min(start + chunk_words, len(words))
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        start = end - overlap_words
    return chunks


class BertDocumentAssistant:
    """Load BERT models once and expose transparent, task-specific methods."""

    def __init__(self) -> None:
        self.device = 0 if torch.cuda.is_available() else -1

        ner_tokenizer = AutoTokenizer.from_pretrained(NER_MODEL_NAME)
        ner_model = AutoModelForTokenClassification.from_pretrained(NER_MODEL_NAME)
        self.ner = pipeline(
            "token-classification",
            model=ner_model,
            tokenizer=ner_tokenizer,
            # "first" labels whole words, so a word is never split into fragments such as "R" + "##udra"
            aggregation_strategy="first",
            device=self.device,
        )

        # Transformers 5 no longer exposes a ``question-answering`` pipeline
        # shortcut. Keep the fine-tuned BERT model and perform extractive QA
        # directly, which also makes the selected source span transparent.
        self.qa_tokenizer = AutoTokenizer.from_pretrained(QA_MODEL_NAME)
        self.qa_model = AutoModelForQuestionAnswering.from_pretrained(QA_MODEL_NAME)
        self.qa_model.eval()
        if self.device == 0:
            self.qa_model.to("cuda")

        self.embedding_tokenizer = AutoTokenizer.from_pretrained(EMBEDDING_MODEL_NAME)
        self.embedding_model = AutoModel.from_pretrained(EMBEDDING_MODEL_NAME)
        self.embedding_model.eval()
        if self.device == 0:
            self.embedding_model.to("cuda")

    def scan_named_entities(self, text: str) -> list[Finding]:
        """Find people, organisations, and locations with BERT NER."""
        if not text.strip():
            return []

        label_map = {"PER": "Person", "ORG": "Organisation", "LOC": "Location"}
        findings: list[Finding] = []
        seen: set[tuple[str, str]] = set()

        for chunk in split_for_qa(text):
            for entity in self.ner(chunk):
                raw_label = entity["entity_group"]
                if raw_label not in label_map:
                    continue
                value = entity["word"].strip()
                key = (value.lower(), raw_label)
                # Skip single characters, leftover word pieces and low-confidence guesses
                if len(value) < 2 or value.startswith("##") or entity["score"] < MIN_ENTITY_SCORE:
                    continue
                if key not in seen:
                    findings.append(Finding(value, label_map[raw_label], float(entity["score"])))
                    seen.add(key)
        return findings

    @staticmethod
    def scan_structured_values(text: str) -> list[Finding]:
        """Find common fixed-format values that a word-label model may miss."""
        patterns = {
            "Email": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
            "Phone": r"(?<!\w)(?:\+?\d{1,3}[ .-]?)?(?:\(?\d{2,4}\)?[ .-]?)?\d{3,4}[ .-]\d{4}(?!\w)",
            "Employee ID": r"\b(?:EMP|EMPLOYEE)[ -]?\d{4,8}\b",
            "Date": r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{1,2},?\s+\d{4})\b",
        }
        findings: list[Finding] = []
        for category, pattern in patterns.items():
            findings.extend(Finding(match.group(), category) for match in re.finditer(pattern, text, re.IGNORECASE))
        return findings

    def scan_document(self, text: str) -> list[Finding]:
        """Combine BERT NER with transparent matching for fixed-format values."""
        findings = self.scan_named_entities(text) + self.scan_structured_values(text)
        unique: dict[tuple[str, str], Finding] = {}
        for finding in findings:
            unique[(finding.text.lower(), finding.category)] = finding
        return sorted(unique.values(), key=lambda item: (item.category, item.text.lower()))

    @staticmethod
    def redact_text(text: str, findings: Iterable[Finding]) -> str:
        """Create a preview with each detected value replaced by its category."""
        # One pass over the text with the longest values first, matching whole words only,
        # so a short value never replaces part of a longer one or part of an inserted tag.
        categories = {finding.text.lower(): finding.category.upper() for finding in findings}
        if not categories:
            return text
        alternatives = "|".join(re.escape(value) for value in sorted(categories, key=len, reverse=True))
        pattern = re.compile(rf"(?<!\w)(?:{alternatives})(?!\w)", re.IGNORECASE)
        return pattern.sub(lambda match: f"[{categories[match.group().lower()]}]", text)

    def answer_question(self, question: str, text: str) -> dict[str, object] | None:
        """Answer a question by selecting the strongest extractive QA answer."""
        if not question.strip() or not text.strip():
            return None

        candidates = []
        for chunk in split_for_qa(text):
            result = self._answer_chunk(question, chunk)
            if result is not None:
                candidates.append({**result, "context": chunk})
        if candidates:
            return max(candidates, key=lambda item: item["confidence"])

        # Table rows are flattened during DOCX extraction. If BERT cannot
        # resolve a clear row relationship, recover an exact Name/Title pair
        # from the row as a document-structure fallback.
        return self._answer_table_row(question, text)

    @staticmethod
    def _answer_table_row(question: str, text: str) -> dict[str, object] | None:
        """Resolve simple Organisation, Name, and Title relationships in tables."""
        question_lower = question.lower()
        if not any(term in question_lower for term in ("who", "whose name", "what is the name")):
            return None

        row_pattern = re.compile(
            r"(?:Table\s+\d+:\s*)?(?P<organisation>.+?)\s*(?:\||;)\s*"
            r"Name:\s*(?P<name>.+?)\s*(?:\||;)\s*Title:\s*(?P<title>[^|;\n]+)",
            re.IGNORECASE,
        )
        for match in row_pattern.finditer(text):
            organisation = match.group("organisation").strip()
            name = match.group("name").strip()
            title = match.group("title").strip()
            if organisation.lower() in question_lower and title.lower() in question_lower:
                return {
                    "answer": name,
                    "confidence": 1.0,
                    "context": match.group(0),
                }
        return None

    def _answer_chunk(self, question: str, chunk: str) -> dict[str, object] | None:
        """Select the best BERT answer span from one document chunk."""
        target_device = "cuda" if self.device == 0 else "cpu"
        encoded = self.qa_tokenizer(
            question,
            chunk,
            truncation="only_second",
            max_length=384,
            return_tensors="pt",
        ).to(target_device)

        with torch.no_grad():
            output = self.qa_model(**encoded)

        start_logits = output.start_logits[0].detach().cpu()
        end_logits = output.end_logits[0].detach().cpu()
        input_ids = encoded["input_ids"][0].detach().cpu()
        context_mask = encoded.get("token_type_ids", torch.ones_like(encoded["input_ids"]))[0].detach().cpu().bool()

        # The first token is [CLS]. It represents the model's no-answer score
        # for the SQuAD 2 checkpoint and is excluded from answer candidates.
        null_score = float(start_logits[0] + end_logits[0])
        valid_positions = torch.where(context_mask)[0].tolist()
        if not valid_positions:
            return None

        best_score = float("-inf")
        best_start = 0
        best_end = 0
        for start in valid_positions:
            for end in range(start, min(start + 48, len(input_ids))):
                if not context_mask[end]:
                    break
                score = float(start_logits[start] + end_logits[end])
                if score > best_score:
                    best_score = score
                    best_start = start
                    best_end = end

        # Require the selected span to score above the checkpoint's null span.
        if best_score <= null_score:
            return None
        answer = self.qa_tokenizer.decode(input_ids[best_start : best_end + 1], skip_special_tokens=True).strip()
        if not answer:
            return None

        # This relative value is intended as an evidence-strength indicator,
        # not a calibrated probability or factual guarantee.
        confidence = float(torch.sigmoid(torch.tensor(best_score - null_score)))
        return {"answer": answer, "confidence": confidence}

    def embed(self, texts: list[str], batch_size: int = 12) -> torch.Tensor:
        """Create mean-pooled, L2-normalised base BERT embeddings."""
        if not texts:
            return torch.empty((0, 768))

        batches: list[torch.Tensor] = []
        target_device = "cuda" if self.device == 0 else "cpu"
        with torch.no_grad():
            for start in range(0, len(texts), batch_size):
                encoded = self.embedding_tokenizer(
                    texts[start : start + batch_size],
                    padding=True,
                    truncation=True,
                    max_length=128,
                    return_tensors="pt",
                ).to(target_device)
                output = self.embedding_model(**encoded).last_hidden_state
                mask = encoded["attention_mask"].unsqueeze(-1)
                pooled = (output * mask).sum(dim=1) / mask.sum(dim=1).clamp(min=1)
                batches.append(torch.nn.functional.normalize(pooled, p=2, dim=1).cpu())
        return torch.cat(batches, dim=0)

    def assess_ambiguity(self, sentence: str) -> dict[str, object]:
        """Compare a sentence with clear and ambiguous business-language examples."""
        clear_examples = [
            "The supplier must submit the report by 5:00 PM on 30 June 2026.",
            "The company will pay the invoice within 30 calendar days of receipt.",
            "Either party may terminate this agreement by giving 14 days written notice.",
        ]
        ambiguous_examples = [
            "The supplier should respond promptly when practical.",
            "The company may take reasonable action as needed.",
            "Payment will be made soon after approval.",
            "The parties will cooperate in good faith where appropriate.",
        ]
        vectors = self.embed([sentence] + clear_examples + ambiguous_examples)
        input_vector = vectors[0]
        clear_score = float(torch.mean(vectors[1 : 1 + len(clear_examples)] @ input_vector))
        ambiguous_score = float(torch.mean(vectors[1 + len(clear_examples) :] @ input_vector))

        cues = [
            cue for cue in ["reasonable", "promptly", "soon", "as needed", "where appropriate", "good faith", "may", "practical"]
            if cue in sentence.lower()
        ]
        verdict = "Potentially ambiguous" if ambiguous_score >= clear_score or cues else "More specific"
        return {
            "verdict": verdict,
            "clear_score": clear_score,
            "ambiguous_score": ambiguous_score,
            "cues": cues,
        }

    def extractive_summary(self, text: str, max_sentences: int = 4) -> list[str]:
        """Rank source sentences by closeness to the document embedding centroid."""
        sentences = split_sentences(text)
        if len(sentences) <= max_sentences:
            return sentences
        vectors = self.embed(sentences)
        centroid = torch.nn.functional.normalize(vectors.mean(dim=0), p=2, dim=0)
        scores = vectors @ centroid
        best_indices = sorted(torch.topk(scores, k=max_sentences).indices.tolist())
        return [sentences[index] for index in best_indices]

    def review_clauses(self, text: str) -> list[dict[str, object]]:
        """Compare document chunks with plain-language clause descriptions using BERT embeddings."""
        chunks = split_for_qa(text, chunk_words=90, overlap_words=15)
        if not chunks:
            return []
        labels = list(CLAUSE_DESCRIPTIONS)
        vectors = self.embed(chunks + [CLAUSE_DESCRIPTIONS[label] for label in labels])
        chunk_vectors = vectors[: len(chunks)]
        label_vectors = vectors[len(chunks) :]
        scores = chunk_vectors @ label_vectors.T

        reviews: list[dict[str, object]] = []
        for label_index, label in enumerate(labels):
            best_chunk_index = int(torch.argmax(scores[:, label_index]))
            reviews.append(
                {
                    "category": label,
                    "similarity": float(scores[best_chunk_index, label_index]),
                    "evidence": chunks[best_chunk_index],
                }
            )
        return sorted(reviews, key=lambda item: item["similarity"], reverse=True)
